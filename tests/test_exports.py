import os
import tempfile
import unittest
from datetime import date
from io import BytesIO
from pathlib import Path
from unittest.mock import patch

from src.exporters import create_docx, create_pdf
from src.orchestrator import run_pipeline
from src.schemas import MeetingMetadata
from src.transcription import TranscriptionError, transcribe_audio


class ExportAndAudioTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.protocol = run_pipeline(
            MeetingMetadata(
                title="Hackathon protocol", meeting_date=date(2026, 9, 23),
                participants=["Aigerim", "Timur", "Dana"],
            ),
            demo_mode=True,
        )

    def test_docx_contains_required_protocol_sections(self):
        from docx import Document

        data = create_docx(self.protocol)
        document = Document(BytesIO(data))
        text = "\n".join(paragraph.text for paragraph in document.paragraphs)
        text += "\n" + "\n".join(cell.text for table in document.tables for row in table.rows for cell in row.cells)
        for expected in ("Hackathon protocol", "2026-09-23", "Participants", "Speaker Mapping", "Action Items", "Meeting Summary", "Transcript", "Timur"):
            self.assertIn(expected, text)

    def test_pdf_is_valid_and_has_multiple_pages(self):
        from pypdf import PdfReader

        data = create_pdf(self.protocol)
        reader = PdfReader(BytesIO(data))
        self.assertGreaterEqual(len(reader.pages), 2)
        self.assertEqual(reader.metadata.title, "Meeting Protocol - Hackathon protocol")

    def test_audio_requires_explicit_local_model_path(self):
        with patch.dict(os.environ, {}, clear=True):
            with self.assertRaisesRegex(TranscriptionError, "QURYLTAI_WHISPER_MODEL_PATH"):
                transcribe_audio(b"audio", "meeting.wav")

    def test_audio_transcription_uses_local_model_and_deletes_temporary_file(self):
        class Segment:
            def __init__(self, text):
                self.text = text

        class FakeModel:
            path = None

            def transcribe(self, path, **kwargs):
                self.path = Path(path)
                self.kwargs = kwargs
                self.path.read_bytes()
                return iter([Segment(" First line. "), Segment("Second line.")]), object()

        model = FakeModel()
        with tempfile.TemporaryDirectory() as directory, patch(
            "src.transcription.get_whisper_model_path", return_value=Path(directory),
        ), patch("src.transcription._load_model", return_value=model):
            result = transcribe_audio(b"not-real-audio", "meeting.wav")
        self.assertEqual(result, "First line.\nSecond line.")
        self.assertFalse(model.path.exists())
        self.assertTrue(model.kwargs["vad_filter"])


if __name__ == "__main__":
    unittest.main()
