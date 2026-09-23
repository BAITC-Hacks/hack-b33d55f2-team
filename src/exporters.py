"""In-memory DOCX and PDF protocol exports."""

from html import escape
from io import BytesIO
from pathlib import Path

from src.schemas import Protocol


class ExportError(RuntimeError):
    """A recoverable document export failure."""


def _speaker_names(protocol: Protocol, display_names: dict[str, str] | None) -> dict[str, str]:
    names = {speaker.id: speaker.name for speaker in protocol.speakers}
    if display_names:
        names.update({key: value for key, value in display_names.items() if value.strip()})
    return names


def _assignee_name(protocol: Protocol, names: dict[str, str], assignee: str | None) -> str:
    if not assignee:
        return "Not specified"
    original_to_display = {speaker.name: names[speaker.id] for speaker in protocol.speakers}
    return original_to_display.get(assignee, assignee)


def create_docx(protocol: Protocol, display_names: dict[str, str] | None = None) -> bytes:
    try:
        from docx import Document
        from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn
        from docx.shared import Inches, Pt
    except ImportError as exc:
        raise ExportError("DOCX export requires python-docx. Install the project requirements.") from exc

    names = _speaker_names(protocol, display_names)
    document = Document()
    section = document.sections[0]
    section.top_margin = section.bottom_margin = Inches(0.7)
    section.left_margin = section.right_margin = Inches(0.75)
    styles = document.styles
    styles["Normal"].font.name = "Arial"
    styles["Normal"].font.size = Pt(10)
    for style_name in ("Title", "Heading 1", "Heading 2"):
        styles[style_name].font.name = "Arial"
        styles[style_name].font.color.rgb = None

    title = document.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run("Meeting Protocol")
    document.add_paragraph(f"{protocol.metadata.title}\n{protocol.metadata.meeting_date.isoformat()}")

    document.add_heading("Participants", level=1)
    participants = protocol.metadata.participants or list(names.values())
    document.add_paragraph(", ".join(participants) if participants else "Not specified")

    document.add_heading("Speaker Mapping", level=1)
    mapping = document.add_table(rows=1, cols=2)
    mapping.style = "Table Grid"
    mapping.rows[0].cells[0].text = "Speaker ID"
    mapping.rows[0].cells[1].text = "Name"
    for speaker in protocol.speakers:
        cells = mapping.add_row().cells
        cells[0].text, cells[1].text = speaker.id, names[speaker.id]

    document.add_heading("Action Items", level=1)
    actions = document.add_table(rows=1, cols=4)
    actions.style = "Table Grid"
    for cell, text in zip(actions.rows[0].cells, ("Responsible", "Task", "Deadline", "Status")):
        cell.text = text
    for item in protocol.action_items:
        cells = actions.add_row().cells
        deadline = item.deadline_original or (
            item.deadline_normalized.isoformat() if item.deadline_normalized else "Not specified"
        )
        values = (_assignee_name(protocol, names, item.assignee), item.task, deadline, item.status)
        for cell, value in zip(cells, values):
            cell.text = value

    for table in (mapping, actions):
        table.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
        for row_index, row in enumerate(table.rows):
            for cell in row.cells:
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                for paragraph in cell.paragraphs:
                    paragraph.paragraph_format.space_after = Pt(2)
                    for run in paragraph.runs:
                        run.font.name = "Arial"
                        run.font.size = Pt(9)
                        run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Arial")
                        if row_index == 0:
                            run.bold = True

    document.add_heading("Meeting Summary", level=1)
    document.add_paragraph(protocol.summary.overview)
    sections = (
        ("Main Topics", protocol.summary.main_topics or protocol.summary.discussion_points),
        ("Key Problems", protocol.summary.key_problems),
        ("Decisions", protocol.summary.decisions),
    )
    for heading, items in sections:
        if items:
            document.add_heading(heading, level=2)
            for item in items:
                document.add_paragraph(item, style="List Bullet")

    document.add_heading("Transcript", level=1)
    for segment in protocol.transcript:
        paragraph = document.add_paragraph()
        label = paragraph.add_run(f"{names.get(segment.speaker_id, segment.speaker_id)}: ")
        label.bold = True
        paragraph.add_run(segment.text)

    output = BytesIO()
    document.save(output)
    return output.getvalue()


def _pdf_fonts():
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont

    candidates = [
        (Path(r"C:\Windows\Fonts\arial.ttf"), Path(r"C:\Windows\Fonts\arialbd.ttf")),
        (Path("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"), Path("/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf")),
    ]
    for regular, bold in candidates:
        if regular.exists() and bold.exists():
            pdfmetrics.registerFont(TTFont("Quryltai", str(regular)))
            pdfmetrics.registerFont(TTFont("Quryltai-Bold", str(bold)))
            return "Quryltai", "Quryltai-Bold"
    raise ExportError("PDF export needs a Unicode font such as Arial or DejaVu Sans.")


def create_pdf(protocol: Protocol, display_names: dict[str, str] | None = None) -> bytes:
    try:
        from reportlab.lib import colors
        from reportlab.lib.enums import TA_CENTER
        from reportlab.lib.pagesizes import A4
        from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
        from reportlab.lib.units import mm
        from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
    except ImportError as exc:
        raise ExportError("PDF export requires ReportLab. Install the project requirements.") from exc

    regular, bold = _pdf_fonts()
    names = _speaker_names(protocol, display_names)
    output = BytesIO()
    document = SimpleDocTemplate(
        output, pagesize=A4, rightMargin=16 * mm, leftMargin=16 * mm,
        topMargin=14 * mm, bottomMargin=14 * mm,
        title=f"Meeting Protocol - {protocol.metadata.title}",
    )
    base = getSampleStyleSheet()
    body = ParagraphStyle("Body", parent=base["BodyText"], fontName=regular, fontSize=9.5, leading=13, spaceAfter=5)
    heading = ParagraphStyle("Heading", parent=base["Heading1"], fontName=bold, fontSize=14, leading=17, textColor=colors.black, spaceBefore=10, spaceAfter=6)
    subheading = ParagraphStyle("Subheading", parent=heading, fontSize=11, leading=14, spaceBefore=7)
    title = ParagraphStyle("Title", parent=base["Title"], fontName=bold, fontSize=20, leading=24, alignment=TA_CENTER, textColor=colors.black)
    cell = ParagraphStyle("Cell", parent=body, fontSize=8.5, leading=11, wordWrap="CJK")
    cell_bold = ParagraphStyle("CellBold", parent=cell, fontName=bold)
    story = [
        Paragraph("Meeting Protocol", title),
        Paragraph(escape(protocol.metadata.title), heading),
        Paragraph(f"Date: {protocol.metadata.meeting_date.isoformat()}", body),
        Paragraph("Participants", heading),
        Paragraph(escape(", ".join(protocol.metadata.participants) or "Not specified"), body),
        Paragraph("Speaker Mapping", heading),
    ]

    mapping_data = [[Paragraph("Speaker ID", cell_bold), Paragraph("Name", cell_bold)]]
    mapping_data += [[Paragraph(escape(s.id), cell), Paragraph(escape(names[s.id]), cell)] for s in protocol.speakers]
    mapping = Table(mapping_data, colWidths=[45 * mm, 115 * mm], repeatRows=1)
    mapping.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#DCEAF2")),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#D9D9D9")),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("LEFTPADDING", (0, 0), (-1, -1), 6), ("RIGHTPADDING", (0, 0), (-1, -1), 6),
        ("TOPPADDING", (0, 0), (-1, -1), 5), ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
    ]))
    story += [mapping, Paragraph("Action Items", heading)]

    action_data = [[Paragraph(text, cell_bold) for text in ("Responsible", "Task", "Deadline", "Status")]]
    for item in protocol.action_items:
        deadline = item.deadline_original or (item.deadline_normalized.isoformat() if item.deadline_normalized else "Not specified")
        action_data.append([
            Paragraph(escape(_assignee_name(protocol, names, item.assignee)), cell),
            Paragraph(escape(item.task), cell), Paragraph(escape(deadline), cell),
            Paragraph(escape(item.status), cell),
        ])
    if len(action_data) == 1:
        action_data.append([Paragraph("No action items", cell), "", "", ""])
    actions = Table(action_data, colWidths=[34 * mm, 70 * mm, 38 * mm, 18 * mm], repeatRows=1)
    actions.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#DCEAF2")),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#D9D9D9")),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("LEFTPADDING", (0, 0), (-1, -1), 4), ("RIGHTPADDING", (0, 0), (-1, -1), 4),
        ("TOPPADDING", (0, 0), (-1, -1), 5), ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
    ]))
    story += [actions, Paragraph("Meeting Summary", heading), Paragraph(escape(protocol.summary.overview), body)]
    for section_name, items in (
        ("Main Topics", protocol.summary.main_topics or protocol.summary.discussion_points),
        ("Key Problems", protocol.summary.key_problems), ("Decisions", protocol.summary.decisions),
    ):
        if items:
            story.append(Paragraph(section_name, subheading))
            story.extend(Paragraph(f"- {escape(item)}", body) for item in items)
    story += [PageBreak(), Paragraph("Transcript", heading)]
    for segment in protocol.transcript:
        speaker = escape(names.get(segment.speaker_id, segment.speaker_id))
        story.append(Paragraph(f"<b>{speaker}:</b> {escape(segment.text)}", body))
        story.append(Spacer(1, 1.5 * mm))

    document.build(story)
    return output.getvalue()
